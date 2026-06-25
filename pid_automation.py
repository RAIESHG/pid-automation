#!/usr/bin/env python3
"""
P&ID Data Entry Automation
Extract tags from a PDF (with real coordinates) -> validate -> place into a DXF
on a REVIEW layer -> export a tag list + BOM to Excel.

IMPORTANT: ezdxf reads/writes DXF, not DWG. In AutoCAD, SAVEAS -> DXF before
running this, then SAVEAS -> DWG afterwards (or use the ODA File Converter).

Usage:
    python pid_automation.py --pdf input.pdf --dxf drawing.dxf

Tag patterns can be overridden with --inst-pattern and --line-pattern.
"""

import re
import sys
import argparse
import logging
from pathlib import Path
from collections import defaultdict
from typing import List, Optional

import pdfplumber
import ezdxf
import pdfplumber.page

# Monkey-patch to capture OCG / properties tags from pdfminer.
# Wrapped in try/except so a pdfplumber API change never breaks the import.
try:
    from pdfplumber.page import PDFPageAggregatorWithMarkedContent
    pdfplumber.page.ALL_ATTRS.add("props")

    _orig_begin_tag    = PDFPageAggregatorWithMarkedContent.begin_tag
    _orig_tag_cur_item = PDFPageAggregatorWithMarkedContent.tag_cur_item

    def _patched_begin_tag(self, tag, props=None):
        _orig_begin_tag(self, tag, props)
        self.cur_props = props

    def _patched_tag_cur_item(self):
        _orig_tag_cur_item(self)
        if self.cur_item._objs:
            self.cur_item._objs[-1].props = getattr(self, 'cur_props', None)

    PDFPageAggregatorWithMarkedContent.begin_tag    = _patched_begin_tag
    PDFPageAggregatorWithMarkedContent.tag_cur_item = _patched_tag_cur_item
except Exception:
    pass  # OCG layer names won't be captured, but all other features still work

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

DEFAULT_INST_PATTERN = (
    r'(?:\d{4}-)?(?!PID-|FID-|PFD-)[A-Z]{1,3}-\d{2,6}[A-Z]?(?:/[A-Z]){0,3}'
)
DEFAULT_LINE_PATTERN = (
    r'(?:[\dX]+(?:/\d+)?"-)?[A-Z]{1,3}-[A-Z0-9X]{3,7}-[A-Z0-9X]{3,4}-[A-Z0-9]{2,4}(?:-[A-Z]{1,2})?'
    r'|[A-Z]{1,3}-[\dX]{3,4}[A-Z]?-[\dX]{3,4}-[A-Z]{1,2}-"[A-Z0-9X]'
)
DOC_REF_PATTERN = r'\d{4,5}-(?:PID|FID|PFD)-\d{4,5}[A-Z]?'

COORD_SCALE = 1.0

EQUIPMENT_PREFIXES = {'P', 'V', 'TK', 'T', 'E', 'C', 'K', 'D', 'R', 'F', 'S', 'X',
                      'HE', 'EH', 'SU', 'YC', 'B', 'M', 'DA', 'EQ', 'TR'}

TITLE_BLOCK_PAGE_FRACTION = 0.7

MARKUP_ACTION_PATTERNS = [
    (r'\b(ADD|NEW|INSERT|INCLUDE)\b',                       'add'),
    (r'\b(DELETE|REMOVE|DEL|OMIT|CANCEL)\b',               'delete'),
    (r'\b(CHANGE|MODIFY|UPDATE|REVISE|REPLACE|CORRECT)\b', 'change'),
    (r'\b(CHECK|VERIFY|CONFIRM|VALIDATE)\b',               'verify'),
    (r'\b(MOVE|RELOCATE|SHIFT)\b',                         'move'),
    (r'\b(HOLD|TBD|TBC|FUTURE)\b',                        'hold'),
]

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s  %(levelname)s  %(message)s',
    datefmt='%H:%M:%S',
)
log = logging.getLogger('pid')


def extract_tags(pdf_path, inst_pattern, line_pattern):
    """Extract tags with real (x, y) positions, one record per matched word."""
    log.info('Extracting from %s', pdf_path)
    inst_re = re.compile(inst_pattern)
    line_re = re.compile(line_pattern)
    doc_re = re.compile(DOC_REF_PATTERN)
    tags = []
    skipped_doc = 0

    with pdfplumber.open(pdf_path) as pdf:
        n_pages = len(pdf.pages)
        for page_num, page in enumerate(pdf.pages, 1):
            page_height = page.height
            words = page.extract_words()
            if not words:
                log.warning('Page %d: no extractable text (scanned image?)', page_num)
                continue

            for w in words:
                text = w['text'].strip().strip(':;,.()[]{}"\'')
                if not text:
                    continue
                if doc_re.fullmatch(text):
                    skipped_doc += 1
                    continue
                x  = w['x0'] * COORD_SCALE
                y  = (page_height - w['top'])    * COORD_SCALE
                x1 = w['x1'] * COORD_SCALE
                y0 = (page_height - w['bottom']) * COORD_SCALE

                if line_re.fullmatch(text):
                    ttype = 'line'
                elif inst_re.fullmatch(text):
                    ttype = 'instrument'
                else:
                    continue

                tags.append({
                    'tag': text, 'type': ttype, 'page': page_num,
                    'x': round(x, 2),  'y':  round(y, 2),
                    'x1': round(x1, 2), 'y0': round(y0, 2),
                    'confidence': 0.95,
                })

    if n_pages > 3:
        page_counts = defaultdict(set)
        for t in tags:
            page_counts[t['tag']].add(t['page'])
        threshold = max(2, int(n_pages * TITLE_BLOCK_PAGE_FRACTION))
        boilerplate = {tag for tag, pages in page_counts.items() if len(pages) >= threshold}
        if boilerplate:
            before = len(tags)
            tags = [t for t in tags if t['tag'] not in boilerplate]
            log.info('Dropped %d title-block occurrences across %d unique tags (e.g. %s)',
                     before - len(tags), len(boilerplate), ', '.join(sorted(boilerplate)[:3]))

    log.info('Extracted %d tags across %d page(s) (skipped %d drawing refs)',
             len(tags), n_pages, skipped_doc)
    return tags


def validate(tags):
    """Flag duplicates, format errors, and numbering gaps within each prefix."""
    log.info('Validating %d tags', len(tags))
    seen = set()
    issues = []

    by_prefix = defaultdict(list)
    for t in tags:
        if t['type'] != 'instrument':
            continue
        m = re.match(r'([A-Z]{1,3})-(\d+)', t['tag'])
        if not m:
            continue
        prefix = m.group(1)
        if prefix in EQUIPMENT_PREFIXES:
            continue
        by_prefix[prefix].append(int(m.group(2)))

    gaps = {}
    for prefix, nums in by_prefix.items():
        nums_sorted = sorted(set(nums))
        if len(nums_sorted) >= 2:
            full = set(range(nums_sorted[0], nums_sorted[-1] + 1))
            missing = sorted(full - set(nums_sorted))
            if missing:
                gaps[prefix] = missing

    for t in tags:
        tag = t['tag']
        status = 'ok'
        if tag in seen:
            status = 'duplicate'
            issues.append(f'duplicate: {tag}')
        else:
            seen.add(tag)
        if len(tag) < 2:
            status = 'format_error'
            issues.append(f'format error: {tag!r}')
        t['status'] = status

    ok = sum(1 for t in tags if t['status'] == 'ok')
    dup = sum(1 for t in tags if t['status'] == 'duplicate')
    err = sum(1 for t in tags if t['status'] == 'format_error')
    log.info('Valid: %d   Duplicates: %d   Format errors: %d', ok, dup, err)

    if gaps:
        for prefix, missing in gaps.items():
            preview = ', '.join(f'{prefix}-{n}' for n in missing[:5])
            more = '' if len(missing) <= 5 else f' (+{len(missing) - 5} more)'
            log.info('Numbering gap in %s-*: missing %s%s', prefix, preview, more)

    if issues:
        for i in issues[:5]:
            log.info('  %s', i)
        if len(issues) > 5:
            log.info('  ... and %d more', len(issues) - 5)

    return tags, gaps


def place_in_dxf(dxf_path, tags, out_path, text_height=2.5, extract_lines=False,
                 symbol_layers='', symbol_blocks='', pdf_geometry=None, pdf_texts=None,
                 shx_texts=None, markup=None):
    """Place each valid tag as MTEXT at its real (x, y) on a REVIEW layer,
    copy any extracted lines/symbols to corresponding REVIEW layers,
    draw extracted PDF vector geometry on their respective layers,
    and place all PDF text on their respective layers.
    """
    log.info('Placing tags and matching geometry into %s', dxf_path)
    try:
        doc = ezdxf.readfile(dxf_path)
    except IOError:
        log.error('Cannot open %s (is it a valid DXF?)', dxf_path)
        return 0
    except ezdxf.DXFStructureError:
        log.error('%s is not a valid DXF. Did you pass a DWG? SAVEAS to DXF first.', dxf_path)
        return 0

    msp = doc.modelspace()
    if 'REVIEW' not in doc.layers:
        doc.layers.new(name='REVIEW', dxfattribs={'color': 1})

    placed = 0
    for t in tags:
        if t['status'] != 'ok':
            continue
        try:
            mtext = msp.add_mtext(t['tag'], dxfattribs={'layer': 'REVIEW'})
            mtext.set_location((t['x'], t['y']))
            mtext.dxf.char_height = text_height
            placed += 1
        except Exception as e:
            log.warning('Failed to place tag %s: %s', t['tag'], e)

    if extract_lines:
        if 'REVIEW_LINES' not in doc.layers:
            doc.layers.new(name='REVIEW_LINES', dxfattribs={'color': 4})
        for line in list(msp.query('LINE')):
            try:
                new_line = line.copy_to_layout(msp)
                new_line.dxf.layer = 'REVIEW_LINES'
            except Exception as e:
                log.warning('Failed to copy line to REVIEW_LINES: %s', e)

    if symbol_layers or symbol_blocks:
        if 'REVIEW_SYMBOLS' not in doc.layers:
            doc.layers.new(name='REVIEW_SYMBOLS', dxfattribs={'color': 3})
        layer_patterns = [p.strip() for p in symbol_layers.split(',') if p.strip()] if symbol_layers else []
        block_patterns = [p.strip() for p in symbol_blocks.split(',') if p.strip()] if symbol_blocks else []
        for ins in list(msp.query('INSERT')):
            layer = ins.dxf.layer
            block = ins.dxf.name
            if layer_patterns and not match_patterns(layer, layer_patterns):
                continue
            if block_patterns and not match_patterns(block, block_patterns):
                continue
            try:
                new_ins = ins.copy_to_layout(msp)
                new_ins.dxf.layer = 'REVIEW_SYMBOLS'
            except Exception as e:
                log.warning('Failed to copy symbol to REVIEW_SYMBOLS: %s', e)

    if pdf_geometry:
        for seg in pdf_geometry:
            start_pt, end_pt, layer_name = seg
            if layer_name not in doc.layers:
                doc.layers.new(name=layer_name, dxfattribs={'color': 7})
            try:
                msp.add_line(start_pt, end_pt, dxfattribs={'layer': layer_name})
            except Exception:
                pass

    if pdf_texts:
        for t in pdf_texts:
            text = t['text']
            x, y = t['x'], t['y']
            if t['is_tag']:
                if t['tag_type'] == 'instrument':
                    target_layer = 'ssp_INSTRUMENTATION'
                elif t['tag_type'] == 'line':
                    target_layer = 'ssp_PIPING'
                else:
                    target_layer = t['layer']
            else:
                target_layer = t['layer']
            if target_layer not in doc.layers:
                doc.layers.new(name=target_layer, dxfattribs={'color': 7})
            try:
                mtext = msp.add_mtext(text, dxfattribs={'layer': target_layer})
                mtext.set_location((x, y))
                mtext.dxf.char_height = text_height
            except Exception as e:
                log.warning('Failed to place text %s: %s', text, e)

    # ── SHX text (AutoCAD font annotations, previously invisible) ────────────────
    if shx_texts:
        if 'ssp_SHX_TEXT' not in doc.layers:
            doc.layers.new(name='ssp_SHX_TEXT', dxfattribs={'color': 7})
        for t in shx_texts:
            char_h = max(2.0, min(t.get('height', text_height), text_height * 1.5))
            try:
                mt = msp.add_mtext(t['text'], dxfattribs={'layer': 'ssp_SHX_TEXT'})
                mt.set_location((t['x'], t['y']))
                mt.dxf.char_height = char_h
            except Exception as e:
                log.warning('Failed to place SHX text %r: %s', t['text'], e)

    # ── Review markup ─────────────────────────────────────────────────────────
    if markup:
        clouds = markup.get('revision_clouds', [])
        if clouds:
            if 'REVISION_CLOUDS' not in doc.layers:
                doc.layers.new(name='REVISION_CLOUDS', dxfattribs={'color': 1})
            for cloud in clouds:
                verts = cloud.get('vertices', [])
                if len(verts) >= 2:
                    try:
                        msp.add_lwpolyline(
                            verts, close=True,
                            dxfattribs={'layer': 'REVISION_CLOUDS'})
                    except Exception as e:
                        log.warning('Failed to place revision cloud: %s', e)

        comments = markup.get('comments', [])
        if comments:
            if 'MARKUP_COMMENTS' not in doc.layers:
                doc.layers.new(name='MARKUP_COMMENTS', dxfattribs={'color': 2})
            for cmt in comments:
                try:
                    label = f"[{cmt.get('action_type', 'note').upper()}] {cmt['text']}"
                    mt = msp.add_mtext(label, dxfattribs={'layer': 'MARKUP_COMMENTS'})
                    mt.set_location((cmt['x'], cmt['y']))
                    mt.dxf.char_height = text_height
                except Exception as e:
                    log.warning('Failed to place comment: %s', e)

        redlines = markup.get('redlines', [])
        if redlines:
            if 'REDLINES' not in doc.layers:
                doc.layers.new(name='REDLINES', dxfattribs={'color': 1})
            for rl in redlines:
                for stroke in rl.get('strokes', []):
                    for i in range(len(stroke) - 1):
                        try:
                            msp.add_line(
                                stroke[i], stroke[i + 1],
                                dxfattribs={'layer': 'REDLINES'})
                        except Exception:
                            pass

        stamps = markup.get('stamps', [])
        if stamps:
            if 'MARKUP_STAMPS' not in doc.layers:
                doc.layers.new(name='MARKUP_STAMPS', dxfattribs={'color': 3})
            for stamp in stamps:
                try:
                    mt = msp.add_mtext(
                        f"[STAMP: {stamp['text']}]",
                        dxfattribs={'layer': 'MARKUP_STAMPS'})
                    mt.set_location((stamp['x'], stamp['y']))
                    mt.dxf.char_height = text_height
                except Exception as e:
                    log.warning('Failed to place stamp: %s', e)

    try:
        doc.saveas(out_path)
        log.info('Placed %d tags and copied/drawn matching geometry -> %s', placed, out_path)
    except PermissionError:
        saved = False
        for i in range(1, 10):
            alt_path = out_path.with_name(f"{out_path.stem}_{i}{out_path.suffix}")
            try:
                doc.saveas(alt_path)
                log.info('Output DXF was locked. Saved instead as -> %s', alt_path.name)
                saved = True
                break
            except PermissionError:
                continue
        if not saved:
            log.error('Could not save DXF: file %s and all alternatives are locked.', out_path.name)
            return 0
    return placed


def export_excel(tags, gaps, out_path, lines: Optional[List[dict]] = None,
                 symbols: Optional[List[dict]] = None,
                 shx_texts: Optional[List[dict]] = None,
                 markup: Optional[dict] = None,
                 llm_data: Optional[dict] = None):
    """Tag list sheet + BOM summary sheet, plus optional Lines and Symbols sheets."""
    log.info('Exporting Excel %s', out_path)
    wb = Workbook()
    header_fill = PatternFill('solid', start_color='4472C4', end_color='4472C4')
    header_font = Font(bold=True, color='FFFFFF', name='Arial')
    body_font = Font(name='Arial')

    def style_header(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

    def autofit(ws, cap=60):
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, cap)

    ws = wb.active
    ws.title = 'Tag List'
    ws.append(['Tag', 'Type', 'Page', 'X', 'Y', 'Status'])
    for t in tags:
        ws.append([t['tag'], t['type'], t['page'], t['x'], t['y'], t['status']])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    style_header(ws)
    autofit(ws)

    bom = wb.create_sheet('BOM')
    bom.append(['Type', 'Count', 'Tags'])
    grouped = defaultdict(list)
    for t in tags:
        if t['status'] == 'ok':
            grouped[t['type']].append(t['tag'])
    for ttype, names in sorted(grouped.items()):
        bom.append([ttype, len(names), ', '.join(sorted(names))])
    for row in bom.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    style_header(bom)
    autofit(bom)

    if gaps:
        gs = wb.create_sheet('Numbering Gaps')
        gs.append(['Prefix', 'Missing'])
        for prefix, missing in gaps.items():
            gs.append([prefix, ', '.join(f'{prefix}-{n}' for n in missing)])
        for row in gs.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        style_header(gs)
        autofit(gs)

    if lines:
        ls = wb.create_sheet('Lines')
        ls.append(['Layer', 'Start X', 'Start Y', 'Start Z', 'End X', 'End Y', 'End Z'])
        for ln in lines:
            ls.append([ln['layer'], ln['start'][0], ln['start'][1], ln['start'][2],
                       ln['end'][0], ln['end'][1], ln['end'][2]])
        for row in ls.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        style_header(ls)
        autofit(ls)

    if symbols:
        ss = wb.create_sheet('Symbols')
        ss.append(['Layer', 'Block', 'X', 'Y', 'Z'])
        for sym in symbols:
            ss.append([sym['layer'], sym['block'],
                       sym['insert_at'][0], sym['insert_at'][1], sym['insert_at'][2]])
        for row in ss.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        style_header(ss)
        autofit(ss)

    if shx_texts:
        shx_ws = wb.create_sheet('SHX Text')
        shx_ws.append(['Text', 'X', 'Y', 'Page'])
        for t in shx_texts:
            shx_ws.append([t['text'], t['x'], t['y'], t['page']])
        for row in shx_ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        style_header(shx_ws)
        autofit(shx_ws)

    if markup:
        action_rows = []
        for cmt in markup.get('comments', []):
            action_rows.append({
                'Page': cmt['page'], 'Type': 'Comment',
                'Action': cmt.get('action_type', 'note'),
                'Referenced Tags': ', '.join(e['tag'] for e in cmt.get('entities', [])),
                'Text': cmt['text'], 'Author': cmt.get('author', ''),
                'X': round(cmt['x'], 1), 'Y': round(cmt['y'], 1),
            })
        for cloud in markup.get('revision_clouds', []):
            b = cloud.get('bbox', (0, 0, 0, 0))
            action_rows.append({
                'Page': cloud['page'], 'Type': 'Revision Cloud',
                'Action': 'revision', 'Referenced Tags': '',
                'Text': f'Area ({b[0]:.0f},{b[1]:.0f})→({b[2]:.0f},{b[3]:.0f})',
                'Author': cloud.get('author', ''),
                'X': round(b[0], 1), 'Y': round(b[1], 1),
            })
        for rl in markup.get('redlines', []):
            b = rl.get('bbox', (0, 0, 0, 0))
            action_rows.append({
                'Page': rl['page'], 'Type': 'Redline',
                'Action': 'redline', 'Referenced Tags': '',
                'Text': f'{len(rl.get("strokes", []))} stroke(s)',
                'Author': rl.get('author', ''),
                'X': round(b[0], 1), 'Y': round(b[1], 1),
            })
        for stamp in markup.get('stamps', []):
            action_rows.append({
                'Page': stamp['page'], 'Type': 'Stamp',
                'Action': stamp['text'], 'Referenced Tags': '',
                'Text': stamp['text'], 'Author': stamp.get('author', ''),
                'X': round(stamp['x'], 1), 'Y': round(stamp['y'], 1),
            })
        if action_rows:
            act_ws = wb.create_sheet('Markup Actions')
            hdrs = ['Page', 'Type', 'Action', 'Referenced Tags', 'Text', 'Author', 'X', 'Y']
            act_ws.append(hdrs)
            for r in action_rows:
                act_ws.append([r.get(h, '') for h in hdrs])
            for row in act_ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
            style_header(act_ws)
            autofit(act_ws)

    if llm_data:
        interps = llm_data.get('interpretations', [])
        applied = llm_data.get('applied', [])
        if interps:
            llm_ws = wb.create_sheet('LLM Interpretations')
            llm_ws.append(['#', 'Original Comment', 'Interpretation', 'Confidence',
                           'Can Automate', 'Actions'])
            for r in interps:
                actions_summary = '; '.join(
                    a.get('type', '') + ': ' + (a.get('text') or a.get('description') or a.get('pattern', ''))
                    for a in r.get('actions', [])
                )
                llm_ws.append([
                    r.get('index', ''), r.get('original_text', ''),
                    r.get('interpretation', ''), r.get('confidence', ''),
                    'Yes' if r.get('can_automate') else 'No',
                    actions_summary,
                ])
            for row in llm_ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
            style_header(llm_ws)
            autofit(llm_ws)
        if applied:
            app_ws = wb.create_sheet('LLM Applied Actions')
            app_ws.append(['Comment #', 'Action Type', 'Detail'])
            for a in applied:
                app_ws.append([a.get('comment', ''), a.get('type', ''), a.get('detail', '')])
            for row in app_ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
            style_header(app_ws)
            autofit(app_ws)

    wb.save(out_path)
    log.info('Wrote %s', out_path)


def extract_lines(dxf_path: str) -> List[dict]:
    """Extract LINE entities from a DXF file."""
    try:
        doc = ezdxf.readfile(dxf_path)
    except Exception as e:
        log.error('Failed to read DXF for line extraction: %s', e)
        return []
    lines = []
    for line in doc.modelspace().query('LINE'):
        lines.append({
            'layer': line.dxf.layer,
            'start': (line.dxf.start.x, line.dxf.start.y, line.dxf.start.z),
            'end': (line.dxf.end.x, line.dxf.end.y, line.dxf.end.z),
        })
    return lines


def match_patterns(name: str, patterns: List[str]) -> bool:
    """Return True if name matches any regex in patterns."""
    for pat in patterns:
        if re.search(pat, name):
            return True
    return False


def extract_symbols(dxf_path: str, layer_patterns: List[str] = None,
                    block_patterns: List[str] = None) -> List[dict]:
    """Extract INSERT (block reference) entities from a DXF file."""
    try:
        doc = ezdxf.readfile(dxf_path)
    except Exception as e:
        log.error('Failed to read DXF for symbol extraction: %s', e)
        return []
    symbols = []
    for ins in doc.modelspace().query('INSERT'):
        layer = ins.dxf.layer
        block = ins.dxf.name
        if layer_patterns and not match_patterns(layer, layer_patterns):
            continue
        if block_patterns and not match_patterns(block, block_patterns):
            continue
        symbols.append({
            'layer': layer,
            'block': block,
            'insert_at': (ins.dxf.insert.x, ins.dxf.insert.y, ins.dxf.insert.z),
        })
    return symbols


def interpolate_bezier(p0, p1, p2, p3, steps=4):
    """Interpolate cubic Bezier curve points into linear segments."""
    pts = []
    for i in range(steps + 1):
        t = i / steps
        x = ((1-t)**3 * p0[0] +
             3 * (1-t)**2 * t * p1[0] +
             3 * (1-t) * t**2 * p2[0] +
             t**3 * p3[0])
        y = ((1-t)**3 * p0[1] +
             3 * (1-t)**2 * t * p1[1] +
             3 * (1-t) * t**2 * p2[1] +
             t**3 * p3[1])
        pts.append((x, y))
    segments = []
    for i in range(len(pts) - 1):
        segments.append((pts[i], pts[i+1]))
    return segments


def extract_pdf_geometry(pdf_path: Path) -> List[tuple]:
    """Extract line segments, rectangles, and curves from PDF vector graphics."""
    log.info('Extracting PDF vector geometry from %s', pdf_path)
    geometry = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                page_height = page.height

                def flip_y(y_val):
                    return (page_height - y_val) * COORD_SCALE

                properties = {}
                resources = page.page_obj.resources
                if resources and 'Properties' in resources:
                    props_dict = resources['Properties']
                    for k, v in props_dict.items():
                        resolved = v.resolve() if hasattr(v, 'resolve') else v
                        if isinstance(resolved, dict) and 'Name' in resolved:
                            name_bytes = resolved['Name']
                            if isinstance(name_bytes, bytes):
                                name = pdfplumber.utils.decode_text(name_bytes)
                            else:
                                name = str(name_bytes)
                            properties[k] = name

                def get_layer(item):
                    prop_val = item.get('props')
                    prop_name = None
                    if prop_val:
                        prop_name = prop_val.name if hasattr(prop_val, 'name') else str(prop_val)
                    return properties.get(prop_name, '0')

                def parse_path_ops_with_layer(path_ops: list, layer_name: str) -> List[tuple]:
                    segments = []
                    current_point = None
                    start_point = None
                    for op, *args in path_ops:
                        if op == 'm':
                            current_point = args[0]
                            if start_point is None:
                                start_point = current_point
                        elif op == 'l':
                            to_point = args[0]
                            if current_point:
                                segments.append((
                                    (current_point[0] * COORD_SCALE, flip_y(current_point[1])),
                                    (to_point[0] * COORD_SCALE, flip_y(to_point[1])),
                                    layer_name
                                ))
                            current_point = to_point
                        elif op == 'c':
                            ctrl1, ctrl2, to_point = args
                            if current_point:
                                bezier_segs = interpolate_bezier(
                                    (current_point[0] * COORD_SCALE, flip_y(current_point[1])),
                                    (ctrl1[0] * COORD_SCALE, flip_y(ctrl1[1])),
                                    (ctrl2[0] * COORD_SCALE, flip_y(ctrl2[1])),
                                    (to_point[0] * COORD_SCALE, flip_y(to_point[1])),
                                    steps=4
                                )
                                for s0, s1 in bezier_segs:
                                    segments.append((s0, s1, layer_name))
                            current_point = to_point
                        elif op == 'h':
                            if current_point and start_point and current_point != start_point:
                                segments.append((
                                    (current_point[0] * COORD_SCALE, flip_y(current_point[1])),
                                    (start_point[0] * COORD_SCALE, flip_y(start_point[1])),
                                    layer_name
                                ))
                            current_point = start_point
                    return segments

                for line in page.lines:
                    layer = get_layer(line)
                    if 'path' in line and line['path']:
                        geometry.extend(parse_path_ops_with_layer(line['path'], layer))
                    elif 'pts' in line and len(line['pts']) >= 2:
                        pts = line['pts']
                        for i in range(len(pts) - 1):
                            geometry.append((
                                (pts[i][0] * COORD_SCALE, flip_y(pts[i][1])),
                                (pts[i+1][0] * COORD_SCALE, flip_y(pts[i+1][1])),
                                layer
                            ))
                    else:
                        x0, y0 = line['x0'], line['top']
                        x1, y1 = line['x1'], line['bottom']
                        geometry.append((
                            (x0 * COORD_SCALE, flip_y(y0)),
                            (x1 * COORD_SCALE, flip_y(y1)),
                            layer
                        ))

                for rect in page.rects:
                    layer = get_layer(rect)
                    if 'path' in rect and rect['path']:
                        geometry.extend(parse_path_ops_with_layer(rect['path'], layer))
                    else:
                        x0, y0 = rect['x0'], rect['top']
                        x1, y1 = rect['x1'], rect['bottom']
                        geometry.append(((x0 * COORD_SCALE, flip_y(y0)), (x1 * COORD_SCALE, flip_y(y0)), layer))
                        geometry.append(((x1 * COORD_SCALE, flip_y(y0)), (x1 * COORD_SCALE, flip_y(y1)), layer))
                        geometry.append(((x1 * COORD_SCALE, flip_y(y1)), (x0 * COORD_SCALE, flip_y(y1)), layer))
                        geometry.append(((x0 * COORD_SCALE, flip_y(y1)), (x0 * COORD_SCALE, flip_y(y0)), layer))

                for curve in page.curves:
                    layer = get_layer(curve)
                    if 'path' in curve and curve['path']:
                        geometry.extend(parse_path_ops_with_layer(curve['path'], layer))
                    elif 'pts' in curve and len(curve['pts']) >= 2:
                        pts = curve['pts']
                        for i in range(len(pts) - 1):
                            geometry.append((
                                (pts[i][0] * COORD_SCALE, flip_y(pts[i][1])),
                                (pts[i+1][0] * COORD_SCALE, flip_y(pts[i+1][1])),
                                layer
                            ))
    except Exception as e:
        log.error('Failed to extract geometry from PDF: %s', e)

    log.info('Extracted %d segment(s) of drawing geometry from PDF', len(geometry))
    return geometry


def extract_pdf_texts(pdf_path: Path, inst_pattern: str, line_pattern: str) -> List[dict]:
    """Extract all text words from PDF with coordinates, OCG layers, and tag classification."""
    log.info('Extracting all text from %s', pdf_path)
    inst_re = re.compile(inst_pattern)
    line_re = re.compile(line_pattern)
    doc_re = re.compile(DOC_REF_PATTERN)

    all_texts = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_height = page.height
                words = page.extract_words()
                if not words:
                    continue

                properties = {}
                resources = page.page_obj.resources
                if resources and 'Properties' in resources:
                    props_dict = resources['Properties']
                    for k, v in props_dict.items():
                        resolved = v.resolve() if hasattr(v, 'resolve') else v
                        if isinstance(resolved, dict) and 'Name' in resolved:
                            name_bytes = resolved['Name']
                            if isinstance(name_bytes, bytes):
                                name = pdfplumber.utils.decode_text(name_bytes)
                            else:
                                name = str(name_bytes)
                            properties[k] = name

                for w in words:
                    text = w['text'].strip().strip(':;,.()[]{}"\'')
                    if not text:
                        continue
                    if doc_re.fullmatch(text):
                        continue

                    x = w['x0'] * COORD_SCALE
                    y = (page_height - w['top']) * COORD_SCALE

                    prop_val = None
                    for char in page.chars:
                        if (char['x0'] >= w['x0'] - 1.0 and
                            char['x1'] <= w['x1'] + 1.0 and
                            char['top'] >= w['top'] - 1.0 and
                            char['bottom'] <= w['bottom'] + 1.0):
                            if 'props' in char:
                                prop_val = char['props']
                                break

                    prop_name = prop_val.name if hasattr(prop_val, 'name') else str(prop_val) if prop_val else None
                    layer = properties.get(prop_name, 'ssp_TEXT')
                    if layer == '0':
                        layer = 'ssp_TEXT'

                    tag_type = None
                    if line_re.fullmatch(text):
                        tag_type = 'line'
                    elif inst_re.fullmatch(text):
                        tag_type = 'instrument'

                    all_texts.append({
                        'text': text,
                        'x':  round(x, 2),
                        'y':  round(y, 2),
                        'x1': round(w['x1'] * COORD_SCALE, 2),
                        'y0': round((page_height - w['bottom']) * COORD_SCALE, 2),
                        'layer': layer,
                        'is_tag': tag_type is not None,
                        'tag_type': tag_type,
                        'page': page_num,
                    })
    except Exception as e:
        log.error('Failed to extract PDF texts: %s', e)

    log.info('Extracted %d text elements from PDF', len(all_texts))
    return all_texts


# ── Annotation helpers ────────────────────────────────────────────────────────

def _decode_annot_bytes(raw) -> str:
    """Decode annotation Contents field (UTF-16-BE or Latin-1 bytes, or plain str)."""
    if isinstance(raw, bytes):
        if raw.startswith(b'\xfe\xff'):
            return raw.decode('utf-16-be', errors='replace').lstrip('﻿')
        return raw.decode('latin-1', errors='replace')
    return str(raw) if raw else ''


def _annot_subtype(a: dict) -> str:
    """Return the bare annotation subtype name, e.g. 'Square', 'Ink', 'FreeText'."""
    return str(a['data'].get('Subtype', '')).strip("/'")


def _annot_xy(a: dict, page_height: float):
    """Return (x, y) in DXF coordinates for the top-left of an annotation bbox."""
    return a['x0'] * COORD_SCALE, (page_height - a['top']) * COORD_SCALE


def extract_shx_annotations(pdf_path: Path) -> List[dict]:
    """Return text records for AutoCAD SHX annotations (Square subtype, title 'AutoCAD SHX Text').

    AutoCAD exports SHX-font text as PDF Square annotations because SHX glyphs
    cannot embed as real PDF text streams.  pdfplumber's extract_words() misses
    them entirely, so they must be read from page.annots.
    """
    log.info('Extracting AutoCAD SHX text annotations from %s', pdf_path)
    results: List[dict] = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                ph = page.height
                for a in page.annots:
                    if a.get('title') != 'AutoCAD SHX Text':
                        continue
                    text = a.get('contents') or _decode_annot_bytes(
                        a['data'].get('Contents', b''))
                    text = text.strip()
                    if not text:
                        continue
                    x, y = _annot_xy(a, ph)
                    h = abs(a['bottom'] - a['top']) * COORD_SCALE
                    results.append({
                        'text': text,
                        'x': round(x, 2),
                        'y': round(y, 2),
                        'height': round(h, 2),
                        'page': page_num,
                    })
    except Exception as e:
        log.error('Failed to extract SHX annotations: %s', e)
    log.info('Extracted %d SHX text annotation(s)', len(results))
    return results


def parse_markup_action(text: str) -> dict:
    """Classify a markup comment and extract any P&ID tag references it mentions."""
    upper = text.upper()
    action_type = 'note'
    for pattern, atype in MARKUP_ACTION_PATTERNS:
        if re.search(pattern, upper):
            action_type = atype
            break
    inst_re = re.compile(DEFAULT_INST_PATTERN)
    line_re = re.compile(DEFAULT_LINE_PATTERN)
    entities: list = []
    for m in inst_re.finditer(text):
        entities.append({'tag': m.group(), 'type': 'instrument'})
    for m in line_re.finditer(text):
        entities.append({'tag': m.group(), 'type': 'line'})
    return {'action_type': action_type, 'entities': entities}


def _raw_verts_to_dxf(flat) -> list:
    """Convert a flat [x1,y1,x2,y2,...] array of raw PDF coords to DXF (x,y) pairs."""
    flat = list(flat)
    return [(flat[i] * COORD_SCALE, flat[i + 1] * COORD_SCALE)
            for i in range(0, len(flat) - 1, 2)]


def extract_markup_annotations(pdf_path: Path) -> dict:
    """Extract review markup from PDF annotations.

    Returns a dict with keys:
      revision_clouds  – PolyLine / Polygon annotations (typically drawn as clouds in review tools)
      comments         – FreeText / Text / Popup with parsed action type and referenced tags
      redlines         – Ink (freehand) annotations, split into individual strokes
      stamps           – Stamp annotations (e.g. 'FOR REVIEW', 'APPROVED')
      highlights       – Highlight / Underline / StrikeOut annotations
    """
    log.info('Extracting markup annotations from %s', pdf_path)
    markup: dict = {
        'revision_clouds': [],
        'comments': [],
        'redlines': [],
        'stamps': [],
        'highlights': [],
    }
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                ph = page.height
                for a in page.annots:
                    if a.get('title') == 'AutoCAD SHX Text':
                        continue
                    sub = _annot_subtype(a)
                    data = a['data']
                    x, y = _annot_xy(a, ph)
                    bbox = (
                        a['x0'] * COORD_SCALE,
                        (ph - a['bottom']) * COORD_SCALE,
                        a['x1'] * COORD_SCALE,
                        (ph - a['top']) * COORD_SCALE,
                    )

                    if sub in ('PolyLine', 'Polygon'):
                        raw = data.get('Vertices', [])
                        try:
                            verts = _raw_verts_to_dxf(raw) if raw else []
                        except Exception:
                            verts = []
                        if verts:
                            markup['revision_clouds'].append({
                                'page': page_num, 'bbox': bbox,
                                'vertices': verts,
                                'author': a.get('title', ''),
                                'subtype': sub,
                            })
                        else:
                            # No vertex data — treat as a comment
                            text = (a.get('contents') or
                                    _decode_annot_bytes(data.get('Contents', b''))).strip()
                            if text:
                                parsed = parse_markup_action(text)
                                markup['comments'].append(
                                    {'page': page_num, 'x': x, 'y': y, 'bbox': bbox,
                                     'text': text, 'author': a.get('title', ''), **parsed})

                    elif sub == 'Ink':
                        ink_list = data.get('InkList', [])
                        strokes = []
                        for stroke_raw in ink_list:
                            try:
                                strokes.append(_raw_verts_to_dxf(stroke_raw))
                            except Exception:
                                pass
                        if strokes:
                            markup['redlines'].append({
                                'page': page_num, 'bbox': bbox,
                                'strokes': strokes,
                                'author': a.get('title', ''),
                            })

                    elif sub in ('FreeText', 'Text', 'Popup'):
                        text = (a.get('contents') or
                                _decode_annot_bytes(data.get('Contents', b''))).strip()
                        if text:
                            parsed = parse_markup_action(text)
                            markup['comments'].append(
                                {'page': page_num, 'x': x, 'y': y, 'bbox': bbox,
                                 'text': text, 'author': a.get('title', ''), **parsed})

                    elif sub == 'Stamp':
                        text = (a.get('contents') or
                                str(data.get('Name', 'STAMP'))).strip()
                        markup['stamps'].append({
                            'page': page_num, 'x': x, 'y': y, 'bbox': bbox,
                            'text': text, 'author': a.get('title', ''),
                        })

                    elif sub in ('Highlight', 'Underline', 'StrikeOut', 'Squiggly'):
                        text = (a.get('contents') or '').strip()
                        markup['highlights'].append({
                            'page': page_num, 'x': x, 'y': y, 'bbox': bbox,
                            'subtype': sub, 'text': text,
                            'author': a.get('title', ''),
                        })

    except Exception as e:
        log.error('Failed to extract markup annotations: %s', e)

    log.info(
        'Markup found — clouds:%d  comments:%d  redlines:%d  stamps:%d  highlights:%d',
        len(markup['revision_clouds']), len(markup['comments']),
        len(markup['redlines']), len(markup['stamps']), len(markup['highlights']),
    )
    return markup


def main():
    ap = argparse.ArgumentParser(description='P&ID data entry automation')
    ap.add_argument('--pdf', required=True, help='input PDF (text-based)')
    ap.add_argument('--dxf', required=True, help='input DXF (SAVEAS from DWG first)')
    ap.add_argument('--inst-pattern', default=DEFAULT_INST_PATTERN)
    ap.add_argument('--line-pattern', default=DEFAULT_LINE_PATTERN)
    ap.add_argument('--text-height', type=float, default=5.0)
    ap.add_argument('--extract-lines', action='store_true')
    ap.add_argument('--symbol-layers', default='')
    ap.add_argument('--symbol-blocks', default='')
    args = ap.parse_args()

    pdf_path, dxf_path = Path(args.pdf), Path(args.dxf)
    if not pdf_path.exists():
        log.error('PDF not found: %s', pdf_path)
        sys.exit(1)
    if not dxf_path.exists():
        log.error('DXF not found: %s', dxf_path)
        sys.exit(1)

    tags = extract_tags(pdf_path, args.inst_pattern, args.line_pattern)
    if not tags:
        log.error('No tags extracted. Check the PDF is text-based and patterns match.')
        sys.exit(1)

    tags, gaps = validate(tags)

    lines_data = []
    symbols_data = []
    if args.extract_lines:
        lines_data = extract_lines(dxf_path)
    if args.symbol_layers or args.symbol_blocks:
        layer_patterns = [p.strip() for p in args.symbol_layers.split(',') if p.strip()]
        block_patterns = [p.strip() for p in args.symbol_blocks.split(',') if p.strip()]
        symbols_data = extract_symbols(dxf_path, layer_patterns=layer_patterns, block_patterns=block_patterns)

    pdf_geometry = extract_pdf_geometry(pdf_path)
    pdf_texts    = extract_pdf_texts(pdf_path, args.inst_pattern, args.line_pattern)
    shx_texts    = extract_shx_annotations(pdf_path)
    markup       = extract_markup_annotations(pdf_path)

    out_dxf = dxf_path.with_name(dxf_path.stem + '_tags_REVIEW.dxf')
    place_in_dxf(
        dxf_path, tags, out_dxf,
        text_height=args.text_height,
        extract_lines=args.extract_lines,
        symbol_layers=args.symbol_layers,
        symbol_blocks=args.symbol_blocks,
        pdf_geometry=pdf_geometry,
        pdf_texts=pdf_texts,
        shx_texts=shx_texts,
        markup=markup,
    )

    out_xlsx = dxf_path.with_name(dxf_path.stem + '_tag_list.xlsx')
    export_excel(tags, gaps, out_xlsx, lines=lines_data, symbols=symbols_data,
                 shx_texts=shx_texts, markup=markup)

    log.info('Done. Open %s in AutoCAD, verify the REVIEW layer, then commit.', out_dxf.name)
    log.info('Excel deliverable: %s', out_xlsx.name)


if __name__ == '__main__':
    main()
