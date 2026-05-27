"""Run pid_automation across every PDF in the folder. Summarize results."""
import subprocess
import sys
from pathlib import Path
from openpyxl import load_workbook

PDFS = sorted(p for p in Path('.').glob('*.pdf') if p.name != 'sample.pdf')
DXF = 'sample.dxf'
OUT_DIR = Path('out')
OUT_DIR.mkdir(exist_ok=True)

print(f'{"PDF":<60} {"inst":>6} {"line":>6} {"dup":>5} {"err":>5} {"placed":>7}')
print('-' * 95)
grand = {'inst': 0, 'line': 0, 'dup': 0, 'err': 0, 'placed': 0}

for pdf in PDFS:
    base = pdf.stem.replace(' ', '_').replace('(', '').replace(')', '')[:40]
    safe_pdf = OUT_DIR / f'{base}.dxf'
    # copy reference dxf so each run writes its own outputs
    subprocess.run(['cp', DXF, str(safe_pdf)], check=True)
    result = subprocess.run([
        sys.executable, 'pid_automation.py',
        '--pdf', str(pdf), '--dxf', str(safe_pdf),
    ], capture_output=True, text=True)

    xlsx = safe_pdf.with_name(safe_pdf.stem + '_tag_list.xlsx')
    inst = line = dup = err = placed = 0
    if xlsx.exists():
        wb = load_workbook(xlsx)
        ws = wb['Tag List']
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            ttype = row[1]
            status = row[5]
            if status == 'ok':
                if ttype == 'instrument':
                    inst += 1
                elif ttype == 'line':
                    line += 1
                placed += 1
            elif status == 'duplicate':
                dup += 1
            elif status == 'format_error':
                err += 1

    print(f'{pdf.name[:60]:<60} {inst:>6} {line:>6} {dup:>5} {err:>5} {placed:>7}')
    grand['inst'] += inst
    grand['line'] += line
    grand['dup'] += dup
    grand['err'] += err
    grand['placed'] += placed

print('-' * 95)
print(f'{"TOTAL":<60} {grand["inst"]:>6} {grand["line"]:>6} {grand["dup"]:>5} {grand["err"]:>5} {grand["placed"]:>7}')
