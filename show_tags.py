"""Print Tag List sheet from the generated xlsx."""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1])
ws = wb['Tag List']
widths = [12, 12, 6, 10, 10, 8]
for row in ws.iter_rows(values_only=True):
    print('  '.join(str(c if c is not None else '').ljust(w) for c, w in zip(row, widths)))
